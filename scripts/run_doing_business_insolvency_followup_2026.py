#!/usr/bin/env python3
"""Frozen, country-holdout Doing Business insolvency follow-up.

--preflight reads registration and identifier keys only. --lock may read selected
values only after the first registration commit; each case receives a separate
source-lock commit. --run fits all three cells after all locks; --verify exactly
reconstructs the source slice and every saved result packet.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import subprocess
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd
import pyarrow.parquet as pq
import scipy.stats as scipy_stats
import statsmodels.formula.api as smf
import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RUNNER = "scripts/run_doing_business_insolvency_followup_2026.py"
PLAN = "engine/audits/doing_business_insolvency_followup_prereg_2026-10-10.md"
ROSTER = "data/pinned_evidence/doing_business_insolvency_followup_2026_key_roster.csv"
ROSTER_SHA = "528960e23fa54a820482c23ea3707f77fe433faea9a677c90a6225a0b8014482"
SLICE = "data/pinned_evidence/doing_business_insolvency_followup_2026_source_slice.csv"
SUMMARY = "engine/audits/doing_business_insolvency_followup_results_2026-10-10.md"
FAMILY = "doing_business_insolvency_later_outcomes_holdout_2026_10_10"
DB_URL = "https://archive.doingbusiness.org/content/dam/doingBusiness/excel/db-2021/Historical-Data--DB04-DB20-.xlsx"
DB_SHA = "95bce841cd858fcf32f74fda67bc786881de4e91e1db522c78fc93f15abb265e"
DB_SHEET = "DB21 Data"
WDI = "data/vintages/world_bank_wdi/NE.GDI.FTOT.ZS@2026-04-30T140059Z.parquet"
WDI_SHA = "fee6f66739512936ab17e82c4b7b788dc3030620ba30a6e5671e972a3a341a21"
WDI_URL = "https://api.worldbank.org/v2/country/all/indicator/NE.GDI.FTOT.ZS?format=json"
WDI_MANIFEST = "data/manifests/fetch_run_2026-04-30T140119Z.yaml"
PWT = "data/vintages/pwt/rtfpna@2026-05-05T195242Z.parquet"
PWT_SHA = "c919c4e3f1e766651cdf4ef3e252b99164fd2dacd234be830d98879df1cf800a"
PWT_URL = "https://dataverse.nl/api/access/datafile/354095"
PWT_MANIFEST = "data/manifests/fetch_run_2026-05-05T200056Z.yaml"
YEARS = (2015, 2016, 2017, 2018)
P_MAX = 0.01 / 3
FAMILY_SIZE = 3
OLD = frozenset("USA GBR DEU FRA ITA ESP NLD BEL AUT SWE NOR DNK FIN IRL CHE PRT GRC JPN CAN AUS NZL KOR ISR EST POL CZE HUN CHL MEX ZAF TUR BRA ARG COL THA MYS IDN PHL IND CHN".split())
CASE = {
    "db_insolvency_framework_later_investment_2026": {
        "column": 83, "x": "framework_index", "y": "investment_share", "sign": "+",
        "x_scale": 1.0, "x_bounds": (0.0, 16.0), "y_bounds": (0.0, 100.0),
        "min_effect": 0.25, "min_material": 1.0, "min_gap": 2.0,
        "source": "world_bank_doing_business:strength_of_insolvency_framework_index",
    },
    "db_insolvency_recovery_later_investment_2026": {
        "column": 82, "x": "recovery_cents", "y": "investment_share", "sign": "+",
        "x_scale": 10.0, "x_bounds": (0.0, 100.0), "y_bounds": (0.0, 100.0),
        "min_effect": 0.50, "min_material": 5.0, "min_gap": 15.0,
        "source": "world_bank_doing_business:insolvency_recovery_rate",
    },
    "db_insolvency_time_later_tfp_2026": {
        "column": 80, "x": "resolution_years", "y": "tfp_growth", "sign": "-",
        "x_scale": 1.0, "x_bounds": (0.0, 30.0), "y_bounds": (-30.0, 30.0),
        "min_effect": 0.20, "min_material": 0.25, "min_gap": 1.0,
        "source": "world_bank_doing_business:insolvency_time_years",
    },
}
MEMBERS = sorted(CASE)
SLICE_COLUMNS = ("country_iso3", "db_year", "baseline_year", "endpoint_year",
                 "framework_index", "recovery_cents", "resolution_years",
                 "investment_share_endpoint", "tfp_baseline", "tfp_endpoint")
LIMITATIONS = (
    "Publisher-revised Doing Business data describe a standardized hypothetical hotel insolvency "
    "in the largest business city, not actual countrywide proceedings. The World Bank discontinued "
    "Doing Business after integrity concerns. WDI fixed capital formation includes public and "
    "private investment; PWT TFP is an estimated macro residual. Temporal ordering is retrospective, "
    "not an ex-ante forecast. Country and year fixed effects do not establish causality, isolate "
    "bankruptcy reform, compare bailouts, or repair the archived post-shock claim."
)


def digest(body: bytes) -> str:
    return hashlib.sha256(body).hexdigest()


def file_sha(path: Path) -> str:
    return digest(path.read_bytes())


def git_bytes(revision: str, path: str) -> bytes:
    result = subprocess.run(["git", "show", f"{revision}:{path}"], cwd=ROOT, capture_output=True)
    if result.returncode:
        raise ValueError(f"missing committed file {revision}:{path}")
    return result.stdout


def first_addition(path: str) -> str:
    result = subprocess.run(["git", "log", "--diff-filter=A", "--reverse", "--format=%H", "HEAD", "--", path],
                            cwd=ROOT, capture_output=True, text=True, check=True)
    commits = result.stdout.strip().splitlines()
    if len(commits) != 1:
        raise ValueError(f"expected one first-addition commit: {path}")
    return commits[0]


def ancestor(older: str, newer: str) -> bool:
    return subprocess.run(["git", "merge-base", "--is-ancestor", older, newer], cwd=ROOT).returncode == 0


def spec_rel(hid: str) -> str:
    return f"hypotheses/regulatory/{hid}.yaml"


def lock_rel(hid: str) -> str:
    return f"data/pinned_evidence/{hid}.source.json"


def steelman_rel(hid: str) -> str:
    return f"hypotheses/steelman/{hid}.md"


def run_rel(hid: str) -> str:
    return f"engine/runs/{hid}"


def registration_paths() -> list[str]:
    return [PLAN, ROSTER, RUNNER, *(spec_rel(h) for h in MEMBERS),
            *(steelman_rel(h) for h in MEMBERS)]


def roster_rows() -> list[tuple[str, int, int, int]]:
    with (ROOT / ROSTER).open(newline="") as stream:
        reader = csv.DictReader(stream)
        if reader.fieldnames != ["country_iso3", "db_year", "baseline_year", "endpoint_year"]:
            raise ValueError("fixed roster header changed")
        rows = [(r["country_iso3"], int(r["db_year"]), int(r["baseline_year"]),
                 int(r["endpoint_year"])) for r in reader]
    if (file_sha(ROOT / ROSTER) != ROSTER_SHA or len(rows) != 512
            or rows != sorted(set(rows)) or len({r[0] for r in rows}) != 128
            or set(r[1] for r in rows) != set(YEARS)
            or any(c in OLD | {"AFG"} or b != t - 2 or e != t + 1
                   for c, t, b, e in rows)):
        raise ValueError("fixed 512-key country holdout changed")
    return rows


def registered() -> tuple[str, dict[str, dict], list[tuple[str, int, int, int]]]:
    paths = registration_paths()
    commits = {first_addition(p) for p in paths}
    if len(commits) != 1:
        raise ValueError("all family registration files must be first-added together")
    commit = next(iter(commits))
    for path in paths:
        body = (ROOT / path).read_bytes()
        if git_bytes(commit, path) != body or git_bytes("HEAD", path) != body:
            raise ValueError(f"registration edited after first commit: {path}")
    roster = roster_rows()
    specs = {}
    for hid, case in CASE.items():
        spec = yaml.safe_load((ROOT / spec_rel(hid)).read_text())
        thr = spec["falsification"]["threshold"]
        if (spec.get("hypothesis_id") != hid or spec.get("status") != "pre_registered"
                or spec.get("claim_direction") != case["sign"]
                or spec.get("evidence_type") != "associational"
                or spec["sample"]["period"] != [2015, 2019]
                or spec["estimator"]["template"] != "panel_fe"
                or spec["estimator"]["fixed_effects"] != ["country", "db_year"]
                or spec["estimator"]["clustering"] != "country"
                or spec["variables"]["treatment"][0]["source"] != case["source"]
                or thr["family_id"] != FAMILY or thr["family_size"] != FAMILY_SIZE
                or not math.isclose(float(thr["p_max_two_sided_bonferroni"]), P_MAX)
                or thr["fixed_key_roster_sha256"] != ROSTER_SHA
                or thr["source_key_roster_sha256"] != ROSTER_SHA
                or thr["source_workbook_sha256"] != DB_SHA
                or thr["wdi_vintage_sha256"] != WDI_SHA
                or thr["pwt_vintage_sha256"] != PWT_SHA
                or thr["source_lock_path"] != lock_rel(hid)
                or thr["expected_beta_sign"] != case["sign"]
                or thr["meaningful_beta_min_abs"] != case["min_effect"]
                or thr["material_within_country_range_raw_units"] != case["min_material"]
                or thr["min_p90_p10_exposure_gap_raw_units"] != case["min_gap"]
                or any(spec.get(k) != [] for k in ("covers_claims", "linked_hypotheses", "linked_conditions"))):
            raise ValueError(f"frozen specification differs from runner: {hid}")
        specs[hid] = spec
    return commit, specs, roster


def source_files(source_root: Path, workbook: Path) -> dict[str, Path]:
    paths = {"workbook": workbook, "wdi": source_root / WDI, "pwt": source_root / PWT}
    for name, expected in (("workbook", DB_SHA), ("wdi", WDI_SHA), ("pwt", PWT_SHA)):
        path = paths[name]
        if not path.is_file() or file_sha(path) != expected:
            raise ValueError(f"missing or changed exact source: {name}")
    return paths


def workbook_keys(workbook: Path) -> set[tuple[str, int]]:
    sheet = load_workbook(workbook, read_only=True, data_only=True)[DB_SHEET]
    keys: list[tuple[str, int]] = []
    for row in sheet.iter_rows(min_row=6, min_col=1, max_col=5, values_only=True):
        if row[0] and row[4]:
            keys.append((str(row[0]).strip(), int(row[4])))
    if len(keys) != len(set(keys)):
        raise ValueError("duplicate publisher economy/DB-year key")
    return set(keys)


def parquet_keys(path: Path) -> set[tuple[str, int]]:
    rows = pq.read_table(path, columns=["country_iso3", "year"]).to_pylist()
    keys = [(str(r["country_iso3"]), int(r["year"])) for r in rows]
    if len(keys) != len(set(keys)):
        raise ValueError(f"duplicate country/year key in {path.name}")
    return set(keys)


def preflight(source_root: Path, workbook: Path) -> dict:
    commit, _, roster = registered()
    paths = source_files(source_root, workbook)  # hashes only
    db = workbook_keys(paths["workbook"])
    wdi = parquet_keys(paths["wdi"])
    pwt = parquet_keys(paths["pwt"])
    for country, year, baseline, endpoint in roster:
        if ((country, year) not in db or (country, baseline) not in wdi
                or (country, endpoint) not in wdi or (country, baseline) not in pwt
                or (country, endpoint) not in pwt):
            raise ValueError("registered key absent from source key sets")
    return {"registration_commit": commit, "fixed_rows": 512, "countries": 128,
            "db_years": list(YEARS), "source_hashes_match": True,
            "selected_magnitudes_read": False}


def value_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (float, int)) and not isinstance(value, bool):
        return repr(value)
    return str(value).strip()


def parquet_values(path: Path) -> dict[tuple[str, int], object]:
    rows = pq.read_table(path, columns=["country_iso3", "year", "value"]).to_pylist()
    result: dict[tuple[str, int], object] = {}
    for row in rows:
        key = (str(row["country_iso3"]), int(row["year"]))
        if key in result:
            raise ValueError(f"duplicate country/year value key in {path.name}")
        result[key] = row["value"]
    return result


def source_slice_bytes(source_root: Path, workbook: Path, roster: list[tuple[str, int, int, int]]) -> bytes:
    paths = source_files(source_root, workbook)
    wanted = {(c, t) for c, t, _, _ in roster}
    sheet = load_workbook(paths["workbook"], read_only=True, data_only=True)[DB_SHEET]
    header = next(sheet.iter_rows(min_row=5, max_row=5, values_only=True))
    expected = {0: "Country code", 4: "DB year", 80: "Time (years)",
                82: "Recovery rate (cents on the dollar)",
                83: "Strength of insolvency framework index (0-16)"}
    if any(header[i] != label for i, label in expected.items()):
        raise ValueError("publisher workbook sheet/column schema changed")
    db: dict[tuple[str, int], tuple[object, object, object]] = {}
    for row in sheet.iter_rows(min_row=6, values_only=True):
        if not row[0] or not row[4]:
            continue
        key = (str(row[0]).strip(), int(row[4]))
        if key not in wanted:
            continue
        if key in db:
            raise ValueError(f"duplicate selected publisher key {key}")
        db[key] = (row[83], row[82], row[80])
    if set(db) != wanted:
        raise ValueError("publisher selected keys do not equal roster")
    wdi = parquet_values(paths["wdi"])
    pwt = parquet_values(paths["pwt"])
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\n")
    writer.writerow(SLICE_COLUMNS)
    for country, year, baseline, endpoint in roster:
        if ((country, endpoint) not in wdi or (country, baseline) not in pwt
                or (country, endpoint) not in pwt):
            raise ValueError("matched selected outcome key absent")
        framework, recovery, duration = db[(country, year)]
        writer.writerow([country, year, baseline, endpoint, value_text(framework),
                         value_text(recovery), value_text(duration),
                         value_text(wdi[(country, endpoint)]),
                         value_text(pwt[(country, baseline)]),
                         value_text(pwt[(country, endpoint)])])
    return out.getvalue().encode()


def source_lock(hid: str, reg_commit: str, body: bytes) -> dict:
    return {
        "hypothesis_id": hid, "family_id": FAMILY, "family_size": FAMILY_SIZE,
        "registration_commit": reg_commit, "source_sha256": DB_SHA,
        "key_roster_path": ROSTER, "key_roster_sha256": ROSTER_SHA,
        "source_slice_path": SLICE, "source_slice_sha256": digest(body),
        "source_files": {
            "doing_business": {"url": DB_URL, "sha256": DB_SHA,
                               "sheet": DB_SHEET, "license": "CC-BY-4.0"},
            "wdi_gfcf_share": {"path": WDI, "sha256": WDI_SHA,
                               "manifest": WDI_MANIFEST, "url": WDI_URL},
            "pwt_tfp": {"path": PWT, "sha256": PWT_SHA,
                        "manifest": PWT_MANIFEST, "url": PWT_URL},
        },
        "source_note": "Retrospectively corrected publisher workbook; fixed country holdout; no causal credit",
    }


def make_lock(source_root: Path, workbook: Path, hid: str) -> dict:
    reg_commit, _, roster = registered()  # commit must precede selected-value read
    if hid not in CASE:
        raise ValueError("unregistered family member")
    target = ROOT / lock_rel(hid)
    if target.exists():
        raise FileExistsError(f"source lock already exists: {hid}")
    body = source_slice_bytes(source_root, workbook, roster)
    slice_path = ROOT / SLICE
    if slice_path.exists():
        if slice_path.read_bytes() != body:
            raise ValueError("existing frozen source slice differs")
    else:
        slice_path.write_bytes(body)
    lock = source_lock(hid, reg_commit, body)
    target.write_text(json.dumps(lock, indent=2, sort_keys=True) + "\n")
    return {"hypothesis_id": hid, "source_lock_path": lock_rel(hid),
            "source_slice_sha256": digest(body), "source_rows": 512}


def custody(source_root: Path, workbook: Path) -> tuple[str, list[tuple[str, int, int, int]], dict[str, str], bytes]:
    reg_commit, _, roster = registered()
    raw_slice = source_slice_bytes(source_root, workbook, roster)
    if (ROOT / SLICE).read_bytes() != raw_slice:
        raise ValueError("exact source slice replay mismatch")
    slice_commit = first_addition(SLICE)
    if not ancestor(reg_commit, slice_commit) or reg_commit == slice_commit:
        raise ValueError("source slice does not follow registration")
    lock_commits: dict[str, str] = {}
    for hid in MEMBERS:
        rel = lock_rel(hid)
        commit = first_addition(rel)
        if (commit == reg_commit or not ancestor(reg_commit, commit)
                or not ancestor(slice_commit, commit)
                or git_bytes(commit, rel) != (ROOT / rel).read_bytes()
                or git_bytes("HEAD", rel) != (ROOT / rel).read_bytes()):
            raise ValueError(f"source lock chronology or immutability failed: {hid}")
        lock = json.loads((ROOT / rel).read_text())
        if lock != source_lock(hid, reg_commit, raw_slice):
            raise ValueError(f"source lock content differs: {hid}")
        lock_commits[hid] = commit
    if git_bytes("HEAD", SLICE) != raw_slice:
        raise ValueError("source slice changed after source lock")
    return reg_commit, roster, lock_commits, raw_slice


def number(value: object) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return float("nan")
    return result if math.isfinite(result) else float("nan")


def frame_for(hid: str, selected: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    case = CASE[hid]
    x_name = case["x"]
    frame = selected[["country_iso3", "db_year", x_name, "investment_share_endpoint",
                      "tfp_baseline", "tfp_endpoint"]].copy()
    frame["x_raw"] = pd.to_numeric(frame[x_name], errors="coerce")
    frame["x"] = frame["x_raw"] / case["x_scale"]
    if case["y"] == "investment_share":
        frame["y"] = pd.to_numeric(frame["investment_share_endpoint"], errors="coerce")
    else:
        baseline = pd.to_numeric(frame["tfp_baseline"], errors="coerce")
        endpoint = pd.to_numeric(frame["tfp_endpoint"], errors="coerce")
        frame["y"] = 100.0 * np.log(endpoint / baseline) / 3.0
        frame.loc[(baseline <= 0) | (endpoint <= 0) | (baseline > 100) | (endpoint > 100), "y"] = np.nan
    xmin, xmax = case["x_bounds"]
    ymin, ymax = case["y_bounds"]
    finite = np.isfinite(frame["x_raw"].to_numpy(dtype=float)) & np.isfinite(frame["y"].to_numpy(dtype=float))
    valid = finite & frame["x_raw"].between(xmin, xmax).to_numpy() & frame["y"].between(ymin, ymax).to_numpy()
    complete = frame.loc[valid, ["country_iso3", "db_year", "x_raw", "x", "y"]].copy()
    counts = complete.groupby("country_iso3").size()
    complete = complete[complete["country_iso3"].isin(counts[counts >= 3].index)].copy()
    complete.sort_values(["country_iso3", "db_year"], inplace=True)
    stats = {"roster_rows": 512, "raw_complete_rows": int(valid.sum()),
             "complete_rows": len(complete), "countries": int(complete["country_iso3"].nunique()),
             "db_years": sorted(int(y) for y in complete["db_year"].unique()),
             "excluded_due_to_missing_or_bounds": int(512 - valid.sum()),
             "excluded_due_to_short_country_history": int(valid.sum() - len(complete))}
    if not complete.empty:
        spans = complete.groupby("country_iso3")["x_raw"].agg(lambda s: float(s.max() - s.min()))
        stats["material_countries"] = int((spans >= case["min_material"]).sum())
        stats["exposure_p90_p10_gap_raw"] = float(np.percentile(complete["x_raw"], 90) -
                                                np.percentile(complete["x_raw"], 10))
    else:
        stats["material_countries"] = 0
        stats["exposure_p90_p10_gap_raw"] = None
    return complete, stats


def fit(frame: pd.DataFrame) -> tuple[dict | None, str | None]:
    if frame.empty or frame["x"].nunique() < 2:
        return None, "no treatment variation"
    try:
        model = smf.ols("y ~ x + C(country_iso3) + C(db_year)", data=frame)
        if np.linalg.matrix_rank(model.exog) != model.exog.shape[1]:
            return None, "rank-deficient country/DB-year model"
        result = model.fit(cov_type="cluster", cov_kwds={"groups": frame["country_iso3"],
                                                       "use_correction": True}, use_t=True)
        beta = float(result.params["x"])
        se = float(result.bse["x"])
        if not math.isfinite(beta) or not math.isfinite(se) or se <= 0:
            return None, "nonfinite or zero country-cluster uncertainty"
        n_countries = int(frame["country_iso3"].nunique())
        p = float(2 * scipy_stats.t.sf(abs(beta / se), n_countries - 1))
        residual = smf.ols("x ~ C(country_iso3) + C(db_year)", data=frame).fit().resid.to_numpy()
        sq = residual * residual
        total = float(sq.sum())
        if not math.isfinite(total) or total <= 1e-12:
            return None, "no residualized treatment variance"
        by_country: dict[str, float] = defaultdict(float)
        by_year: dict[int, float] = defaultdict(float)
        for country, year, s in zip(frame["country_iso3"], frame["db_year"], sq):
            by_country[str(country)] += float(s)
            by_year[int(year)] += float(s)
        country_share = max(by_country.values()) / total
        year_share = max(by_year.values()) / total
        if country_share > 0.20 or year_share > 0.45:
            return None, "residualized exposure concentration cap failed"
        return {"coefficient": beta, "std_error": se, "p_value_two_sided": p,
                "n_rows": len(frame), "n_countries": n_countries,
                "df_inference": n_countries - 1, "r_squared": float(result.rsquared),
                "top_country_residualized_exposure_share": country_share,
                "top_db_year_residualized_exposure_share": year_share,
                "method": "country/DB-year FE OLS; country CR1; two-sided Student t(G-1)"}, None
    except (ValueError, np.linalg.LinAlgError) as error:
        return None, f"fit failure: {type(error).__name__}"


def evaluate(hid: str, selected: pd.DataFrame) -> dict:
    case = CASE[hid]
    frame, sample = frame_for(hid, selected)
    gate_failures = []
    if (sample["complete_rows"] < 300 or sample["countries"] < 85
            or sample["db_years"] != list(YEARS)):
        gate_failures.append("complete_case_coverage")
    if (sample["material_countries"] < 20
            or sample["exposure_p90_p10_gap_raw"] is None
            or sample["exposure_p90_p10_gap_raw"] < case["min_gap"]):
        gate_failures.append("material_exposure")
    pre = frame[frame["db_year"] <= 2017].copy()
    sample["sensitivity_rows"] = len(pre)
    sample["sensitivity_countries"] = int(pre["country_iso3"].nunique())
    if len(pre) < 200 or sample["sensitivity_countries"] < 70:
        gate_failures.append("sensitivity_coverage")
    primary = sensitivity = None
    if not gate_failures:
        primary, failure = fit(frame)
        if failure:
            gate_failures.append("primary_" + failure.replace(" ", "_"))
        sensitivity, failure = fit(pre)
        if failure:
            gate_failures.append("sensitivity_" + failure.replace(" ", "_"))
    if gate_failures:
        return {"verdict": "INCONCLUSIVE", "reason": "; ".join(gate_failures),
                "gate_failures": gate_failures, "sample": sample,
                "primary": primary, "sensitivity": sensitivity}
    directional = primary["coefficient"] if case["sign"] == "+" else -primary["coefficient"]
    sensitivity_directional = sensitivity["coefficient"] if case["sign"] == "+" else -sensitivity["coefficient"]
    if (directional >= case["min_effect"] and primary["p_value_two_sided"] <= P_MAX
            and sensitivity_directional >= 0):
        verdict, reason = "SUPPORTED", "predicted sign, practical magnitude, family cutoff and sensitivity sign pass"
    elif (directional <= -case["min_effect"] and primary["p_value_two_sided"] <= P_MAX
            and sensitivity_directional <= 0):
        verdict, reason = "REFUTED", "opposite sign, practical magnitude, family cutoff and sensitivity sign pass"
    else:
        verdict, reason = "PARTIAL", "estimable association does not meet every frozen directional gate"
    return {"verdict": verdict, "reason": reason, "gate_failures": [],
            "sample": sample, "primary": primary, "sensitivity": sensitivity}


def artifact_bytes(hid: str, result: dict, timestamp: str,
                   reg_commit: str, lock_commit: str) -> dict[str, bytes]:
    case = CASE[hid]
    source_list = [
        {"publisher": "World Bank Doing Business", "path": "publisher historical workbook",
         "url": DB_URL, "sha256": DB_SHA, "license": "CC-BY-4.0"},
        {"publisher": "World Bank WDI", "path": WDI, "url": WDI_URL, "sha256": WDI_SHA,
         "manifest": WDI_MANIFEST},
        {"publisher": "Penn World Table", "path": PWT, "url": PWT_URL, "sha256": PWT_SHA,
         "manifest": PWT_MANIFEST},
    ]
    diagnostic = {"hypothesis_id": hid, "family_id": FAMILY, "family_size": FAMILY_SIZE,
                  "registration_commit": reg_commit, "source_lock_commit": lock_commit,
                  "source_sha256": DB_SHA, "source_key_roster_sha256": ROSTER_SHA,
                  "source_lock": lock_rel(hid), "source_slice": SLICE,
                  "runner": RUNNER, "run_utc": timestamp, "template": "panel_fe",
                  "result": result, "data_status": {"variables_missing": [], "input_vintages": source_list},
                  "limitations": LIMITATIONS}
    diagnostic_body = (json.dumps(diagnostic, indent=2, sort_keys=True, allow_nan=False) + "\n").encode()
    manifest = {"hypothesis_id": hid, "status": result["verdict"], "runner": RUNNER,
                "preregistration_commit": reg_commit,
                "source": {"source_sha256": DB_SHA, "source_lock_commit": lock_commit,
                           "source_lock_manifest": lock_rel(hid), "source_slice_sha256": file_sha(ROOT / SLICE)},
                "vintages": source_list,
                "family_correction": {"id": FAMILY, "size": 3, "alpha": 0.01, "p_max": P_MAX}}
    packet = {"packet_version": 1,
              "hypothesis": {"id": hid, "path": spec_rel(hid), "topic": "regulatory"},
              "preregistration": {"spec_commit": reg_commit, "source_lock_commit": lock_commit,
                                  "key_roster": ROSTER, "key_roster_sha256": ROSTER_SHA},
              "family": {"id": FAMILY, "size": 3, "members": MEMBERS},
              "source": {"source_sha256": DB_SHA, "source_lock_manifest": lock_rel(hid),
                         "source_lock_commit": lock_commit,
                         "key_roster_sha256": ROSTER_SHA, "input_vintages": source_list},
              "result": result,
              "reproduction": {"runner": RUNNER,
                               "command": "python3 " + RUNNER + " --verify --source-root <pinned-cache-root> --workbook <publisher-xlsx>",
                               "result_sha256": digest(diagnostic_body)},
              "limitations": LIMITATIONS,
              "credit": {"causal": 0, "policy_effect": 0, "school": 0, "position": 0, "featured": 0}}
    primary = result["primary"]
    if primary:
        finding = (f"**{result['verdict']}.** Beta {primary['coefficient']:+.6g} "
                   f"(country-cluster SE {primary['std_error']:.6g}; two-sided p "
                   f"{primary['p_value_two_sided']:.6g}; {primary['n_rows']} country-years, "
                   f"{primary['n_countries']} countries). {result['reason']}.")
    else:
        finding = f"**INCONCLUSIVE.** {result['reason']}. No directional coefficient is graded."
    card = (f"# Doing Business insolvency holdout: {case['x']} → {case['y']}\n\n"
            f"{finding}\n\nThe original 40-country bankruptcy sample and Afghanistan are excluded. "
            f"DB2015-18 indicators precede later macro endpoints; all three correlated tests "
            f"share a two-sided 0.01/3 cutoff. {LIMITATIONS} "
            "No causal, policy, school, position or featured credit.\n")
    replication = ("#!/usr/bin/env python3\nimport os, subprocess, sys\n"
                   "source_root=os.environ.get('IESET_SOURCE_ROOT')\n"
                   "workbook=os.environ.get('IESET_DB_WORKBOOK')\n"
                   "if not source_root or not workbook: raise SystemExit('set IESET_SOURCE_ROOT and IESET_DB_WORKBOOK')\n"
                   f"raise SystemExit(subprocess.call([sys.executable,'{RUNNER}','--verify','--source-root',source_root,'--workbook',workbook]))\n")
    chart = {"hypothesis_id": hid, "verdict": result["verdict"],
             "coefficient": primary.get("coefficient") if primary else None,
             "p_value_two_sided": primary.get("p_value_two_sided") if primary else None,
             "n_rows": result["sample"]["complete_rows"],
             "db_years": list(YEARS)}
    return {"diagnostics.json": diagnostic_body,
            "manifest.yaml": yaml.safe_dump(manifest, sort_keys=False, allow_unicode=True).encode(),
            "evidence_packet.yaml": yaml.safe_dump(packet, sort_keys=False, allow_unicode=True).encode(),
            "hypothesis.yaml": (ROOT / spec_rel(hid)).read_bytes(),
            "result_card.md": card.encode(), "replication.py": replication.encode(),
            "chart_data.json": (json.dumps(chart, indent=2, sort_keys=True) + "\n").encode()}


def summary_bytes(results: dict[str, dict]) -> bytes:
    lines = ["# Doing Business insolvency holdout — measured results", "",
             "Three value-blind registered cells share one fixed country holdout and a two-sided 0.01/3 cutoff. The original post-shock bankruptcy/bailout claim is not graded here.", "",
             "| Cell | Verdict | Complete rows | Countries | Reason |",
             "| --- | --- | ---: | ---: | --- |"]
    for hid in MEMBERS:
        r = results[hid]
        lines.append(f"| `{hid}` | {r['verdict']} | {r['sample']['complete_rows']} | {r['sample']['countries']} | {r['reason']} |")
    lines += ["", LIMITATIONS,
              "", "All cells have zero causal, policy, school, position and featured credit. An INCONCLUSIVE cell receives no public conversion credit. Exact replay requires the hash-pinned publisher workbook and WDI/PWT source files.", ""]
    return "\n".join(lines).encode()


def run(source_root: Path, workbook: Path, verify: bool) -> dict:
    reg_commit, _, lock_commits, _ = custody(source_root, workbook)
    selected = pd.read_csv(ROOT / SLICE, dtype={"country_iso3": str, "db_year": int})
    if tuple(selected.columns) != SLICE_COLUMNS or len(selected) != 512:
        raise ValueError("selected source slice columns/rows changed")
    results = {}
    for hid in MEMBERS:
        result = evaluate(hid, selected)
        results[hid] = result
        directory = ROOT / run_rel(hid)
        if verify:
            if not directory.is_dir():
                raise FileNotFoundError(str(directory))
            saved = json.loads((directory / "diagnostics.json").read_text())
            timestamp = saved["run_utc"]
        else:
            if directory.exists():
                raise FileExistsError(str(directory))
            timestamp = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
        produced = artifact_bytes(hid, result, timestamp, reg_commit, lock_commits[hid])
        if verify:
            changed = [name for name, body in produced.items()
                       if not (directory / name).is_file() or (directory / name).read_bytes() != body]
            if changed:
                raise ValueError(f"exact result packet replay mismatch {hid}: {changed}")
        else:
            directory.mkdir(parents=True)
            for name, body in produced.items():
                (directory / name).write_bytes(body)
    summary = summary_bytes(results)
    if verify:
        if (ROOT / SUMMARY).read_bytes() != summary:
            raise ValueError("exact family summary replay mismatch")
        commits = {first_addition(f"{run_rel(hid)}/diagnostics.json") for hid in MEMBERS}
        if len(commits) != 1 or any(not ancestor(lock_commits[hid], next(iter(commits))) for hid in MEMBERS):
            raise ValueError("family result commit does not follow all source locks")
    else:
        (ROOT / SUMMARY).write_bytes(summary)
    return {hid: {"verdict": result["verdict"], "rows": result["sample"]["complete_rows"],
                  "countries": result["sample"]["countries"]} for hid, result in results.items()}


def main() -> int:
    parser = argparse.ArgumentParser()
    flags = parser.add_mutually_exclusive_group(required=True)
    for flag in ("preflight", "lock", "run", "verify"):
        flags.add_argument(f"--{flag}", action="store_true")
    parser.add_argument("--source-root", type=Path)
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--case", choices=MEMBERS)
    args = parser.parse_args()
    if not args.source_root or not args.workbook:
        parser.error("--source-root and --workbook are required")
    if args.lock and not args.case:
        parser.error("--case is required for --lock")
    if args.preflight:
        output = preflight(args.source_root, args.workbook)
    elif args.lock:
        output = make_lock(args.source_root, args.workbook, args.case)
    else:
        output = run(args.source_root, args.workbook, args.verify)
    print(json.dumps(output, indent=2, sort_keys=True, allow_nan=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
