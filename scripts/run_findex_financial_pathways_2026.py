#!/usr/bin/env python3
"""Frozen, source-locked Global Findex financial-pathways family.

The registration commit contains this runner, six specs, their steelmen, the
preregistration and a key-only roster. Each case gets a later source-lock
commit. Only --run and --verify read selected indicator values or null masks.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import subprocess
from collections import Counter, defaultdict
from pathlib import Path

import numpy as np
import pandas as pd
import patsy
import statsmodels.api as sm
import yaml
from jsonschema import Draft202012Validator
from openpyxl import load_workbook
from scipy.stats import t as student_t

ROOT = Path(__file__).resolve().parents[1]
FAMILY = "findex_financial_pathways_2026_10_10"
RUNNER = "scripts/run_findex_financial_pathways_2026.py"
PREREG = "engine/audits/findex_financial_pathways_prereg_2026-10-10.md"
SUMMARY = "engine/audits/findex_financial_pathways_results_2026-10-10.md"
ROSTER = "data/pinned_evidence/findex_financial_pathways_2026_key_roster.csv"
ROSTER_SHA = "767cdc617b9d210a4ddaa088ece47bbb22a1f73a573a5fb748485554314ad777"
RAW = "data/vintages/world_bank_findex/global_findex_2025@2026-10-10.xlsx"
RAW_SHA = "b6c991f8e46bcbd4b6423cc962d5b2fa030e3db99832b9016bb4f400f57c772f"
SOURCE_URL = "https://thedocs.worldbank.org/en/doc/be6615202d1f08a25855c8ac2d615122-0050012025/related/GlobalFindexDatabase2025.xlsx"
CATALOG_URL = "https://www.worldbank.org/en/publication/globalfindex/download-data"
FAMILY_ALPHA = 0.01
P_MAX = FAMILY_ALPHA / 6
MIN_BETA_ABS = 0.5
MIN_PAIRS = 100
MIN_COUNTRIES = 50
MIN_MATERIAL_COUNTRIES = 20
MATERIAL_CHANGE_PP = 5.0
MAX_COUNTRY_SHARE = 0.20
MAX_TOP_ONE_PERCENT_SHARE = 0.20
MAX_TRANSITION_SHARE = 0.60
ANCHORS = (("2017", "2021"), ("2021", "2024"))
INCOMES = {"Low income", "Lower middle income", "Upper middle income"}
CASES = (
    ("distance_account", "regulatory", "fin11a.s", "all", "all",
     "account.t.d", "all", "all", "-", "physical bank distance versus account uptake"),
    ("cost_poor_account", "distribution", "fin11b.s", "all", "all",
     "account.t.d", "income", "poorest 40%", "-", "service cost versus poor-adult account uptake"),
    ("documentation_women_account", "distribution", "fin11c.s", "all", "all",
     "account.t.d", "gender", "women", "-", "documentation barriers versus women's account uptake"),
    ("mobile_formal_saving", "growth", "mobileaccount.t.d", "all", "all",
     "fin17a.17a1.d", "all", "all", "+", "mobile-money access versus formal saving"),
    ("digital_wages_emergency_funds", "distribution", "fin32.acc.s", "all", "all",
     "fin24aP", "all", "all", "+", "account-paid wages versus emergency funds"),
    ("digital_payments_business_borrowing", "growth", "g20.any", "all", "all",
     "fin22e", "all", "all", "+", "digital payments versus business-purpose borrowing"),
)
LIMITATIONS = [
    "Repeated Findex country surveys do not follow the same adults over time; country averages are ecological estimates.",
    "COVID-era interview modes and later mode changes, questionnaire differences, income and employment composition, and common trends can affect survey-wave changes.",
    "The 2017 and 2021 exposure periods are observational; economy and transition effects do not identify a digital-finance or banking-policy effect.",
    "All six tests share one source and overlapping respondents or constructs, so six records are not independent replications. The family correction is local to these six tests.",
    "Business-purpose borrowing includes informal sources; account ownership and formal saving include mobile-money accounts by definition.",
    "No causal, policy, school, position, or featured credit is permitted.",
]
CORRECTION = {"disposition": "descriptive_calibration_only", "family_id": FAMILY,
              "family_size": 6, "family_alpha": FAMILY_ALPHA,
              "bonferroni_p_max": P_MAX, "school_score_eligible": False,
              "school_score_credit": 0, "policy_score_eligible": False,
              "policy_score_credit": 0, "causal_credit": 0,
              "featured_eligible": False}


def case_id(case: tuple[str, ...]) -> str:
    return f"findex_{case[0]}_forward_2026"


def spec_rel(case: tuple[str, ...]) -> str:
    return f"hypotheses/{case[1]}/{case_id(case)}.yaml"


def steel_rel(case: tuple[str, ...]) -> str:
    return f"hypotheses/steelman/{case_id(case)}.md"


def lock_rel(case: tuple[str, ...]) -> str:
    return f"data/pinned_evidence/{case_id(case)}.source.json"


def run_rel(case: tuple[str, ...]) -> str:
    return f"engine/runs/{case_id(case)}"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def git(*args: str) -> bytes:
    return subprocess.check_output(["git", *args], cwd=ROOT, stderr=subprocess.PIPE)


def first_addition(rel: str) -> str:
    commits = git("log", "--diff-filter=A", "--reverse", "--format=%H", "HEAD", "--", rel).decode().splitlines()
    if len(commits) != 1:
        raise RuntimeError(f"missing or ambiguous first addition: {rel}")
    return commits[0]


def unchanged(rel: str, revision: str) -> None:
    if (ROOT / rel).read_bytes() != git("show", f"{revision}:{rel}"):
        raise RuntimeError(f"frozen bytes changed: {rel}")
    if git("status", "--porcelain", "--", rel).strip():
        raise RuntimeError(f"uncommitted frozen path: {rel}")


def source_path(source_root: Path) -> Path:
    path = source_root / RAW
    if not path.is_file() or sha256(path) != RAW_SHA:
        raise RuntimeError("missing or changed official Global Findex workbook")
    return path


def read_roster() -> list[dict[str, str]]:
    path = ROOT / ROSTER
    if sha256(path) != ROSTER_SHA:
        raise RuntimeError("frozen source-key roster checksum mismatch")
    with path.open(newline="", encoding="utf-8") as file:
        rows = list(csv.DictReader(file))
    names = ["country_iso3", "exposure_year", "outcome_year", "source_region", "source_income_group"]
    if len(rows) != 150 or list(rows[0]) != names:
        raise RuntimeError("frozen source-key roster shape mismatch")
    keys = [(r["country_iso3"], r["exposure_year"]) for r in rows]
    if keys != sorted(set(keys)):
        raise RuntimeError("frozen source-key roster duplicate or unsorted")
    counts = Counter(r["country_iso3"] for r in rows)
    if len(counts) != 75 or any(n != 2 for n in counts.values()):
        raise RuntimeError("frozen roster must contain two transitions for 75 economies")
    if any((r["exposure_year"], r["outcome_year"]) not in ANCHORS
           or r["source_income_group"] not in INCOMES for r in rows):
        raise RuntimeError("frozen roster includes an unregistered transition or income")
    return rows


def read_key_columns_only(raw: Path) -> tuple[set[tuple[str, str, object, object]], dict[tuple[str, str], tuple[str, str]]]:
    workbook = load_workbook(raw, read_only=True, data_only=True)
    sheet = workbook["Data"]
    rows = sheet.iter_rows(min_col=1, max_col=8, values_only=True)
    header = next(rows)
    next(rows)  # English display labels, not source values.
    if header != ("countrynewwb", "codewb", "year", "pop_adult", "regionwb24_hi",
                   "incomegroupwb24", "group", "group2"):
        raise RuntimeError("unexpected source key schema")
    keys: set[tuple[str, str, object, object]] = set()
    details: dict[tuple[str, str], tuple[str, str]] = {}
    for _name, iso, year, _population, region, income, group, subgroup in rows:
        if not isinstance(iso, str) or len(iso) != 3 or str(year) not in {"2017", "2021", "2024"}:
            continue
        key = (iso, str(year), group, subgroup)
        if key in keys:
            raise RuntimeError(f"duplicate source key: {key}")
        keys.add(key)
        if group == "all" and subgroup == "all":
            details[(iso, str(year))] = (region, income)
    workbook.close()
    return keys, details


def preflight(raw: Path, roster: list[dict[str, str]]) -> None:
    keys, details = read_key_columns_only(raw)
    expected: list[tuple[str, str, str, str, str]] = []
    for iso in sorted({k[0] for k in keys}):
        eligible = True
        for anchor, end in ANCHORS:
            detail = details.get((iso, anchor))
            if detail is None or detail[1] not in INCOMES:
                eligible = False
                break
            for year, group, subgroup in ((anchor, "all", "all"),
                                          (anchor, "income", "poorest 40%"),
                                          (end, "income", "poorest 40%"),
                                          (anchor, "gender", "women"),
                                          (end, "gender", "women"),
                                          (end, "all", "all")):
                if (iso, year, group, subgroup) not in keys:
                    eligible = False
                    break
            if not eligible:
                break
        if eligible:
            for anchor, end in ANCHORS:
                region, income = details[(iso, anchor)]
                expected.append((iso, anchor, end, region, income))
    observed = [(r["country_iso3"], r["exposure_year"], r["outcome_year"],
                 r["source_region"], r["source_income_group"]) for r in roster]
    if observed != expected:
        raise RuntimeError("official workbook keys do not reproduce frozen roster")


def threshold(case: tuple[str, ...]) -> dict[str, object]:
    return {"family_id": FAMILY, "family_size": 6, "family_alpha": FAMILY_ALPHA,
            "p_max_two_sided_bonferroni": P_MAX, "expected_beta_sign": case[8],
            "meaningful_beta_min_abs": MIN_BETA_ABS, "min_complete_pairs": MIN_PAIRS,
            "min_complete_countries": MIN_COUNTRIES,
            "min_material_exposure_countries": MIN_MATERIAL_COUNTRIES,
            "material_exposure_change_pp": MATERIAL_CHANGE_PP,
            "max_country_residualized_exposure_share": MAX_COUNTRY_SHARE,
            "max_top_one_percent_residualized_exposure_share": MAX_TOP_ONE_PERCENT_SHARE,
            "max_transition_residualized_exposure_share": MAX_TRANSITION_SHARE,
            "expected_key_pairs": 150, "expected_key_countries": 75,
            "source_key_roster_sha256": ROSTER_SHA,
            "source_lock_path": lock_rel(case), "source_path": RAW,
            "source_sha256": RAW_SHA,
            "school_score_eligible": False, "school_score_credit": 0,
            "policy_score_eligible": False, "policy_score_credit": 0}


def registration_commit() -> str:
    registered = [PREREG, RUNNER, ROSTER]
    for case in CASES:
        registered.extend((spec_rel(case), steel_rel(case)))
    commits = {first_addition(rel) for rel in registered}
    if len(commits) != 1:
        raise RuntimeError("registration files lack one shared first-addition commit")
    revision = commits.pop()
    for rel in registered:
        unchanged(rel, revision)
    schema = json.loads(git("show", f"{revision}:schemas/hypothesis.schema.json"))
    validator = Draft202012Validator(schema)
    for case in CASES:
        spec = yaml.safe_load((ROOT / spec_rel(case)).read_text())
        if not validator.is_valid(spec):
            raise RuntimeError(f"invalid original registered schema: {case_id(case)}")
        if (spec.get("hypothesis_id") != case_id(case)
                or spec.get("falsification", {}).get("threshold") != threshold(case)
                or spec.get("covers_claims") != []
                or spec.get("linked_hypotheses") != []):
            raise RuntimeError(f"frozen registered spec differs from family contract: {case_id(case)}")
        if git("ls-tree", "-r", "--name-only", revision, "--", run_rel(case)).strip():
            raise RuntimeError("registration already contains measured results")
    return revision


def lock_payload(case: tuple[str, ...], registration: str) -> dict[str, object]:
    return {"schema": "ieset-source-lock-v1", "family_id": FAMILY, "family_size": 6,
            "hypothesis_id": case_id(case), "registration_commit": registration,
            "key_roster_sha256": ROSTER_SHA, "source_sha256": RAW_SHA,
            "source": {"publisher": "World Bank Global Findex", "dataset": "Global Findex Database 2025",
                       "path": RAW, "sha256": RAW_SHA, "source_url": SOURCE_URL,
                       "catalog_url": CATALOG_URL, "license": "CC BY 4.0"},
            "measurement_boundary": "Spec and key roster were frozen in an earlier commit; selected indicator values remain unread until this lock is committed."}


def lock_commit(case: tuple[str, ...], registration: str) -> str:
    rel = lock_rel(case)
    revision = first_addition(rel)
    if revision == registration:
        raise RuntimeError("source lock cannot share the registration commit")
    git("merge-base", "--is-ancestor", registration, revision)
    unchanged(rel, revision)
    if json.loads((ROOT / rel).read_text()) != lock_payload(case, registration):
        raise RuntimeError(f"source lock differs from frozen contract: {case_id(case)}")
    return revision


def load_selected_rows(raw: Path, roster: list[dict[str, str]]) -> dict[tuple[str, str, object, object], dict[str, object]]:
    permitted = {r["country_iso3"] for r in roster}
    required_codes = {c[2] for c in CASES} | {c[5] for c in CASES}
    required_groups = {(c[3], c[4]) for c in CASES} | {(c[6], c[7]) for c in CASES}
    workbook = load_workbook(raw, read_only=True, data_only=True)
    sheet = workbook["Data"]
    rows = sheet.iter_rows(values_only=True)
    codes = next(rows)
    next(rows)  # Display labels.
    # The official workbook has repeated empty trailing header cells; require
    # uniqueness only for the selected named indicator codes.
    if any(codes.count(code) != 1 for code in required_codes):
        raise RuntimeError("selected Findex indicator code schema mismatch")
    indices = {code: codes.index(code) for code in required_codes}
    selected: dict[tuple[str, str, object, object], dict[str, object]] = {}
    for row in rows:
        iso, year, group, subgroup = row[1], str(row[2]), row[6], row[7]
        if iso not in permitted or year not in {"2017", "2021", "2024"} or (group, subgroup) not in required_groups:
            continue
        key = (iso, year, group, subgroup)
        if key in selected:
            raise RuntimeError(f"duplicate selected Findex row: {key}")
        selected[key] = {code: row[index] for code, index in indices.items()}
    workbook.close()
    return selected


def percent(value: object) -> float | None:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(result) or result < 0 or result > 1:
        return None
    return result * 100


def measured_case(case: tuple[str, ...], roster: list[dict[str, str]], selected: dict) -> dict[str, object]:
    slug, _topic, xcode, xgroup, xsub, ycode, ygroup, ysub, sign, _label = case
    complete: list[dict[str, object]] = []
    missing = Counter()
    for row in roster:
        iso, anchor, later = row["country_iso3"], row["exposure_year"], row["outcome_year"]
        xrow = selected.get((iso, anchor, xgroup, xsub), {})
        y0row = selected.get((iso, anchor, ygroup, ysub), {})
        y1row = selected.get((iso, later, ygroup, ysub), {})
        x = percent(xrow.get(xcode)); y0 = percent(y0row.get(ycode)); y1 = percent(y1row.get(ycode))
        if x is None:
            missing["exposure_missing_or_out_of_range"] += 1
        if y0 is None:
            missing["baseline_outcome_missing_or_out_of_range"] += 1
        if y1 is None:
            missing["later_outcome_missing_or_out_of_range"] += 1
        if any(v is None for v in (x, y0, y1)):
            continue
        complete.append({"country_iso3": iso, "anchor": anchor,
                         "x": x / 10, "y0": y0 / 10, "dy": y1 - y0,
                         "x_pp": x})
    counts = Counter(row["country_iso3"] for row in complete)
    frame = pd.DataFrame((r for r in complete if counts[r["country_iso3"]] == 2),
                         columns=["country_iso3", "anchor", "x", "y0", "dy", "x_pp"])
    countries = int(frame.country_iso3.nunique()) if not frame.empty else 0
    material = 0
    if not frame.empty:
        material = int(sum(float(g.x_pp.max() - g.x_pp.min()) >= MATERIAL_CHANGE_PP
                           for _, g in frame.groupby("country_iso3")))
    sample: dict[str, object] = {"key_pairs": len(roster), "key_countries": 75,
                                 "complete_before_two_transition_rule": len(complete),
                                 "complete_pairs": len(frame), "countries": countries,
                                 "transitions": sorted(frame.anchor.unique().tolist()) if not frame.empty else [],
                                 "material_exposure_change_countries_5pp": material,
                                 "excluded_cell_counts": dict(sorted(missing.items()))}
    failures = []
    if len(frame) < MIN_PAIRS:
        failures.append("sample_below_100_complete_pairs")
    if countries < MIN_COUNTRIES:
        failures.append("sample_below_50_complete_countries")
    if len(sample["transitions"]) != 2:
        failures.append("both_transitions_not_represented")
    if material < MIN_MATERIAL_COUNTRIES:
        failures.append("material_exposure_change_below_20_countries")
    primary: dict[str, object] | None = None
    if not failures:
        design = patsy.dmatrix("x + y0 + C(country_iso3) + C(anchor)", frame, return_type="dataframe")
        matrix = design.to_numpy(dtype=float)
        if np.linalg.matrix_rank(matrix) != matrix.shape[1]:
            failures.append("design_not_full_rank")
        else:
            z = design.drop(columns=["x"]).to_numpy(dtype=float)
            residual = frame.x.to_numpy(dtype=float) - z @ np.linalg.lstsq(z, frame.x.to_numpy(dtype=float), rcond=None)[0]
            squares = residual ** 2
            total = float(squares.sum())
            if not math.isfinite(total) or total <= 0:
                failures.append("no_residual_exposure")
            else:
                country_shares = defaultdict(float)
                transition_shares = defaultdict(float)
                for country, anchor, square in zip(frame.country_iso3, frame.anchor, squares):
                    country_shares[country] += float(square) / total
                    transition_shares[anchor] += float(square) / total
                top_country = max(country_shares.values())
                top_one_percent = float(sum(sorted(squares, reverse=True)[:max(1, math.ceil(len(squares) * .01))]) / total)
                top_transition = max(transition_shares.values())
                sample["residual_exposure_share"] = {"country": top_country,
                                                     "top_one_percent": top_one_percent,
                                                     "transition": top_transition}
                if top_country > MAX_COUNTRY_SHARE:
                    failures.append("country_exposure_concentration_above_20pct")
                if top_one_percent > MAX_TOP_ONE_PERCENT_SHARE:
                    failures.append("top_one_percent_exposure_concentration_above_20pct")
                if top_transition > MAX_TRANSITION_SHARE:
                    failures.append("transition_exposure_concentration_above_60pct")
        if not failures:
            model = sm.OLS(frame.dy.to_numpy(dtype=float), design).fit(
                cov_type="cluster", cov_kwds={"groups": frame.country_iso3.to_numpy(),
                                              "use_correction": True, "df_correction": True})
            beta = float(model.params["x"])
            se = float(model.bse["x"])
            df = countries - 1
            if not math.isfinite(beta) or not math.isfinite(se) or se <= 0:
                failures.append("nonfinite_or_nonpositive_cluster_uncertainty")
            else:
                p = float(2 * student_t.sf(abs(beta / se), df))
                critical = float(student_t.ppf(1 - P_MAX / 2, df))
                if not math.isfinite(p) or not math.isfinite(critical):
                    failures.append("nonfinite_inference")
                else:
                    primary = {"coefficient_per_10pp_exposure": beta, "std_error": se,
                               "p_value_two_sided": p, "confidence_level": 1 - P_MAX,
                               "confidence_interval": [beta - critical * se, beta + critical * se],
                               "n_pairs": len(frame), "n_countries": countries,
                               "df_inference": df, "design_rank": matrix.shape[1],
                               "method": "economy/transition fixed-effect OLS; baseline outcome; economy-cluster CR1; Student t(G-1)"}
    if failures:
        verdict = "INCONCLUSIVE"
        reason = "one or more frozen source, sample, variation, concentration, or estimator gates did not pass"
        primary = None
    else:
        assert primary is not None
        beta, p = primary["coefficient_per_10pp_exposure"], primary["p_value_two_sided"]
        oriented = beta if sign == "+" else -beta
        if p <= P_MAX and oriented >= MIN_BETA_ABS:
            verdict, reason = "SUPPORTED", "signed descriptive association clears both frozen family and magnitude gates"
        elif p <= P_MAX and oriented <= -MIN_BETA_ABS:
            verdict, reason = "REFUTED", "opposite signed descriptive association clears both frozen family and magnitude gates"
        else:
            verdict, reason = "PARTIAL", "estimable descriptive association does not clear both frozen directional gates"
    return {"verdict": verdict, "reason": reason, "gate_failures": failures,
            "sample": sample, "primary": primary}


def rendered_files(case: tuple[str, ...], result: dict[str, object], registration: str,
                   lock_revision: str) -> dict[str, bytes]:
    hypothesis = case_id(case)
    source = {"publisher": "World Bank Global Findex", "series_id": "Global Findex Database 2025",
              "path": RAW, "sha256": RAW_SHA, "source_url": SOURCE_URL,
              "catalog_url": CATALOG_URL, "license": "CC BY 4.0"}
    diagnostic = {"hypothesis_id": hypothesis, "family_id": FAMILY, "family_size": 6,
                  "registration_commit": registration, "source_lock_commit": lock_revision,
                  "source_sha256": RAW_SHA, "source_key_roster_sha256": ROSTER_SHA,
                  "result": result, "correction": CORRECTION, "limitations": LIMITATIONS}
    diagnostic_bytes = (json.dumps(diagnostic, sort_keys=True, indent=2, ensure_ascii=False) + "\n").encode()
    manifest = {"hypothesis_id": hypothesis, "status": result["verdict"],
                "runner": RUNNER, "preregistration_commit": registration,
                "source_lock": lock_rel(case),
                "source": {"source_sha256": RAW_SHA, "source_lock_commit": lock_revision},
                "vintages": [source], "sample": result["sample"],
                "family_correction": {"size": 6, "alpha": FAMILY_ALPHA, "p_max": P_MAX},
                "correction": CORRECTION}
    packet = {"packet_version": 1,
              "hypothesis": {"id": hypothesis, "path": spec_rel(case), "topic": case[1]},
              "preregistration": {"spec_commit": registration, "source_lock_commit": lock_revision,
                                  "key_roster": ROSTER, "key_roster_sha256": ROSTER_SHA},
              "family": {"id": FAMILY, "size": 6, "members": sorted(case_id(c) for c in CASES)},
              "source": {"source_sha256": RAW_SHA, "source_lock_manifest": lock_rel(case),
                         "input_vintages": [source], "attribution": "World Bank Global Findex",
                         "license": "CC BY 4.0"},
              "result": result, "reproduction": {"runner": RUNNER,
                                                "command": f"python3 {RUNNER} --verify --source-root <pinned-cache-root>",
                                                "result_sha256": hashlib.sha256(diagnostic_bytes).hexdigest()},
              "limitations": LIMITATIONS, "credit": {"causal": 0, "policy_effect": 0,
                                                   "school": 0, "position": 0, "featured": 0},
              "correction": CORRECTION}
    if result["primary"]:
        beta = result["primary"]["coefficient_per_10pp_exposure"]
        p = result["primary"]["p_value_two_sided"]
        figure = f"coefficient {beta:+.6f} outcome points per 10 exposure points; two-sided p={p:.6g}"
    else:
        figure = "no admitted coefficient under the frozen gates"
    card = (f"# {hypothesis} — {result['verdict']}\n\n"
            f"**Registered question:** {case[9]}.\n\n"
            f"**Result:** {figure}. {result['reason'].capitalize()}.\n\n"
            f"**Sample:** {result['sample']['complete_pairs']} complete economy-transition rows, "
            f"{result['sample']['countries']} economies. Gate findings: "
            f"{', '.join(result['gate_failures']) if result['gate_failures'] else 'none'}.\n\n"
            "This is a direct-survey ecological calibration, not an individual effect or policy evaluation. "
            "The six correlated questions share one Bonferroni family; no school, policy, causal, "
            "position or featured credit is assigned.\n")
    wrapper = ("#!/usr/bin/env python3\n"
               "import subprocess, sys\n"
               f"raise SystemExit(subprocess.call([sys.executable, '{RUNNER}', '--verify', "
               "*sys.argv[1:]]))\n")
    return {"diagnostics.json": diagnostic_bytes,
            "manifest.yaml": yaml.safe_dump(manifest, sort_keys=False, allow_unicode=True).encode(),
            "evidence_packet.yaml": yaml.safe_dump(packet, sort_keys=False, allow_unicode=True).encode(),
            "hypothesis.yaml": (ROOT / spec_rel(case)).read_bytes(),
            "result_card.md": card.encode(), "replication.py": wrapper.encode()}


def summary(results: dict[str, dict[str, object]]) -> bytes:
    lines = ["# Global Findex financial pathways — measured results (2026-10-10)", "",
             "The six source-locked, prospectively registered survey questions were measured under their frozen gates. All are descriptive ecological calibrations; no causal, policy, school, position, or featured credit is assigned.", "",
             "| Question | Verdict | Complete pairs | Economies | Association |", "| --- | --- | ---: | ---: | --- |"]
    for case in CASES:
        result = results[case_id(case)]
        primary = result["primary"]
        estimate = (f"β={primary['coefficient_per_10pp_exposure']:+.4f}; p={primary['p_value_two_sided']:.4g}"
                    if primary else "No admitted coefficient")
        lines.append(f"| `{case_id(case)}` | {result['verdict']} | {result['sample']['complete_pairs']} | {result['sample']['countries']} | {estimate} |")
    lines += ["", "Source: World Bank Global Findex Database 2025 country workbook, CC BY 4.0; exact SHA-256 `" + RAW_SHA + "`.",
              "The original data are 2017, 2021, and 2024 waves of repeated adult surveys. The 2021 COVID interview methods and 2024 mode changes limit comparability; questions share source respondents and are not independent experiments.",
              "A family member remains in the six-test Bonferroni denominator even if its frozen coverage, variation, concentration, or estimator gate is inconclusive.", ""]
    return "\n".join(lines).encode()


def execute(mode: str, root: Path) -> None:
    raw = source_path(root)
    roster = read_roster()
    preflight(raw, roster)
    if mode == "preflight":
        print(json.dumps({"source_sha256": RAW_SHA, "key_pairs": 150, "economies": 75,
                          "selected_indicator_values_opened": False}, sort_keys=True))
        return
    registration = registration_commit()
    if mode == "lock":
        raise RuntimeError("--lock requires --case")
    locks = {case_id(c): lock_commit(c, registration) for c in CASES}
    selected = load_selected_rows(raw, roster)
    results = {case_id(c): measured_case(c, roster, selected) for c in CASES}
    generated = {case_id(c): rendered_files(c, results[case_id(c)], registration, locks[case_id(c)])
                 for c in CASES}
    generated["summary"] = {SUMMARY: summary(results)}
    for name, files in generated.items():
        base = ROOT if name == "summary" else ROOT / "engine" / "runs" / name
        if mode == "run":
            base.mkdir(parents=True, exist_ok=True)
        for filename, expected in files.items():
            path = base / filename
            if mode == "run":
                path.write_bytes(expected)
            elif not path.is_file() or path.read_bytes() != expected:
                raise RuntimeError(f"exact replay differs: {path.relative_to(ROOT)}")
    print(json.dumps({"verdicts": dict(sorted(Counter(r["verdict"] for r in results.values()).items())),
                      "cases": {name: {"verdict": r["verdict"], "pairs": r["sample"]["complete_pairs"],
                                       "countries": r["sample"]["countries"]}
                                for name, r in results.items()},
                      "mode": mode}, sort_keys=True))


def main() -> None:
    parser = argparse.ArgumentParser()
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--preflight", action="store_true")
    mode.add_argument("--lock", action="store_true")
    mode.add_argument("--run", action="store_true")
    mode.add_argument("--verify", action="store_true")
    parser.add_argument("--case", choices=[case_id(c) for c in CASES])
    parser.add_argument("--source-root", type=Path, default=ROOT)
    args = parser.parse_args()
    root = args.source_root.resolve()
    if args.preflight:
        execute("preflight", root)
        return
    if args.lock:
        if not args.case:
            parser.error("--lock needs --case")
        raw = source_path(root)
        roster = read_roster()
        preflight(raw, roster)
        registration = registration_commit()
        case = next(c for c in CASES if case_id(c) == args.case)
        rel = lock_rel(case)
        if (ROOT / rel).exists():
            raise RuntimeError("source lock already exists")
        (ROOT / rel).write_text(json.dumps(lock_payload(case, registration),
                                           sort_keys=True, indent=2) + "\n")
        print(json.dumps({"source_lock": rel, "registration_commit": registration,
                          "selected_indicator_values_opened": False}, sort_keys=True))
        return
    execute("run" if args.run else "verify", root)


if __name__ == "__main__":
    main()
